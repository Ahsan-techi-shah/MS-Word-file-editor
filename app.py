import io
import json
import os
import re
from copy import deepcopy

import streamlit as st
from docx import Document
from docx.text.paragraph import Paragraph
from groq import Groq

MODEL = "openai/gpt-oss-20b"
CHUNK_CHARS = 9000
MAX_CHUNKS = 8
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

st.set_page_config(page_title="AI Word Document Editor", page_icon="📝", layout="wide")
st.title("📝 AI Word Document Editor")
st.caption("Upload a Word file → get suggested edits → approve → download.")


# -----------------------------
# Setup
# -----------------------------
def get_client():
    key = os.getenv("GROQ_API_KEY")
    if not key:
        try:
            key = st.secrets["GROQ_API_KEY"]
        except Exception:
            key = None
    if not key:
        raise RuntimeError("GROQ_API_KEY is not configured in secrets.")
    return Groq(api_key=key)


# -----------------------------
# Document reading
# -----------------------------
def extract_units(doc):
    """Return editable units. Each unit keeps direct references to its paragraphs,
    so edits stay correct even after inserting/deleting other paragraphs."""
    units = {}
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            units[f"p{i}"] = {"id": f"p{i}", "kind": "paragraph",
                              "text": p.text, "paras": [p]}

    seen = set()
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                if id(cell._tc) in seen:  # merged cells repeat
                    continue
                seen.add(id(cell._tc))
                text = cell.text.strip()
                if text:
                    uid = f"t{ti}r{ri}c{ci}"
                    units[uid] = {"id": uid, "kind": "cell", "text": text,
                                  "paras": list(cell.paragraphs)}
    return units


def make_chunks(units):
    chunks, current, size = [], [], 0
    for u in units.values():
        line = f"[{u['id']}] {u['text']}"
        if size + len(line) > CHUNK_CHARS and current:
            chunks.append(current)
            current, size = [], 0
        current.append(line)
        size += len(line)
    if current:
        chunks.append(current)
    return chunks[:MAX_CHUNKS], len(chunks) > MAX_CHUNKS


# -----------------------------
# AI
# -----------------------------
SYSTEM_PROMPT = """You are a careful Word document editing assistant.
You receive document text as lines like: [id] text.

Return ONLY valid JSON:
{"suggestions": [ ... ]}

Each suggestion must be one of:
{"op":"find_replace","id":"p3","find":"exact old text","replace":"new text","reason":"why"}
{"op":"replace_text","id":"p3","new_text":"full new sentence/paragraph","reason":"why"}
{"op":"delete_paragraph","id":"p3","reason":"why"}
{"op":"insert_after","id":"p3","new_text":"text to add","reason":"why"}
{"op":"insert_before","id":"p3","new_text":"text to add","reason":"why"}

Rules:
- Use find_replace for numbers, dates, names and small word changes. "find" must be
  copied EXACTLY from the paragraph text.
- Use replace_text only when a whole sentence/paragraph must be rewritten.
- Use only ids that appear in the input. Never invent ids.
- delete_paragraph / insert_* only for ids starting with "p" (not table cells).
- Keep reasons to one short line.
- If nothing needs changing, return {"suggestions": []}.
"""


def call_groq(client, chunk_lines, instruction, focus):
    if instruction.strip():
        task = (
            "Apply this instruction to the document. Find every place it affects "
            f"and propose the exact edits:\n{instruction.strip()}"
        )
    else:
        task = (
            "Review the document and list the edits that are needed. Look for: "
            f"{focus}. Only suggest real problems; do not rewrite things that are fine."
        )

    user = f"{task}\n\nDOCUMENT:\n" + "\n".join(chunk_lines)

    resp = client.chat.completions.create(
        model=MODEL,
        temperature=0.1,
        response_format={"type": "json_object"},
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": user},
        ],
    )
    content = resp.choices[0].message.content.strip()
    content = re.sub(r"^```(?:json)?\s*|\s*```$", "", content)
    return json.loads(content).get("suggestions", [])


# -----------------------------
# Edit helpers (formatting-preserving)
# -----------------------------
def set_paragraph_text(par, text):
    """Replace text but keep the first run's formatting."""
    runs = par.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        par.add_run(text)


def replace_in_paragraph(par, find, repl):
    done = False
    for r in par.runs:  # number/word fix inside a single run keeps all formatting
        if find in r.text:
            r.text = r.text.replace(find, repl)
            done = True
    if not done and find in par.text:  # text split across runs
        set_paragraph_text(par, par.text.replace(find, repl))
        done = True
    return done


def preview_after(unit, s):
    before = unit["text"]
    op = s["op"]
    if op == "find_replace":
        return before.replace(s.get("find", ""), s.get("replace", ""))
    if op == "replace_text":
        return s.get("new_text", "")
    if op == "delete_paragraph":
        return "(paragraph deleted)"
    if op == "insert_after":
        return before + "\n➕ " + s.get("new_text", "")
    if op == "insert_before":
        return "➕ " + s.get("new_text", "") + "\n" + before
    return before


def validate(suggestions, units):
    """Drop anything that can't be applied safely."""
    good = []
    for s in suggestions:
        u = units.get(s.get("id"))
        op = s.get("op")
        if not u:
            continue
        if op == "find_replace":
            find = s.get("find", "")
            if not find or find == s.get("replace", "") or find not in u["text"]:
                continue
        elif op == "replace_text":
            if not s.get("new_text") or s["new_text"].strip() == u["text"].strip():
                continue
        elif op in ("delete_paragraph", "insert_after", "insert_before"):
            if u["kind"] != "paragraph":
                continue
            if op != "delete_paragraph" and not s.get("new_text"):
                continue
        else:
            continue
        good.append(s)
    return good


def apply_suggestions(units, approved):
    applied = 0
    for s in approved:
        u = units[s["id"]]
        paras, op = u["paras"], s["op"]

        if op == "find_replace":
            if any(replace_in_paragraph(p, s["find"], s.get("replace", "")) for p in paras):
                applied += 1
        elif op == "replace_text":
            set_paragraph_text(paras[0], s["new_text"])
            for extra in paras[1:]:
                set_paragraph_text(extra, "")
            applied += 1
        elif op == "delete_paragraph":
            el = paras[0]._element
            el.getparent().remove(el)
            applied += 1
        elif op in ("insert_after", "insert_before"):
            anchor = paras[0]
            new_p = deepcopy(anchor._p)  # copies style of neighbouring paragraph
            (anchor._p.addnext if op == "insert_after" else anchor._p.addprevious)(new_p)
            set_paragraph_text(Paragraph(new_p, anchor._parent), s["new_text"])
            applied += 1
    return applied


# -----------------------------
# UI
# -----------------------------
FOCUS_OPTIONS = {
    "Numbers, amounts, dates & years": "wrong, outdated or inconsistent numbers, amounts, percentages, dates and years",
    "Spelling & grammar": "spelling mistakes and grammar errors",
    "Sentence clarity": "awkward, unclear or wordy sentences that should be reworded",
    "Names, titles & contact details": "inconsistent or wrong names, job titles, addresses, phone numbers and emails",
}

uploaded = st.file_uploader("Upload your Word document (.docx)", type=["docx"])

if not uploaded:
    st.info("Upload a .docx file to begin.")
    st.stop()

file_bytes = uploaded.getvalue()
file_key = f"{uploaded.name}-{len(file_bytes)}"
if st.session_state.get("file_key") != file_key:
    st.session_state.file_key = file_key
    st.session_state.suggestions = None

units = extract_units(Document(io.BytesIO(file_bytes)))
st.success(f"Loaded **{uploaded.name}** — {len(units)} text blocks found.")

col1, col2 = st.columns(2)
with col1:
    instruction = st.text_area(
        "Specific instruction (optional)",
        placeholder="e.g. Change the project cost from 5 million to 7.5 million everywhere, "
                    "and update the year 2024 to 2026.",
        height=120,
    )
with col2:
    focus_labels = st.multiselect(
        "If no instruction, check for:",
        list(FOCUS_OPTIONS),
        default=list(FOCUS_OPTIONS)[:2],
    )

if st.button("🔍 Find needed edits", type="primary"):
    try:
        client = get_client()
        focus = "; ".join(FOCUS_OPTIONS[f] for f in focus_labels) or "any errors"
        chunks, truncated = make_chunks(units)
        found = []
        bar = st.progress(0.0, text="Analysing document...")
        for i, chunk in enumerate(chunks):
            found += call_groq(client, chunk, instruction, focus)
            bar.progress((i + 1) / len(chunks), text=f"Analysed part {i + 1} of {len(chunks)}")
        bar.empty()
        st.session_state.suggestions = validate(found, units)
        if truncated:
            st.warning("Document is long — only the first part was analysed.")
    except Exception as e:
        st.error(f"Error: {e}")

suggestions = st.session_state.get("suggestions")
if suggestions is not None:
    if not suggestions:
        st.info("No edits needed" if not instruction.strip()
                else "Couldn't find anywhere that instruction applies. Try wording it more specifically.")
        st.stop()

    st.subheader(f"{len(suggestions)} suggested edit(s)")
    chosen = []
    for i, s in enumerate(suggestions):
        u = units[s["id"]]
        with st.container(border=True):
            on = st.checkbox(f"**{i + 1}. {s.get('reason', 'Suggested change')}**",
                             value=True, key=f"chk_{file_key}_{i}")
            a, b = st.columns(2)
            a.markdown("**Before**")
            a.text(u["text"])
            b.markdown("**After**")
            b.text(preview_after(u, s))
        if on:
            chosen.append(s)

    if st.button(f"✅ Apply {len(chosen)} selected edit(s)"):
        doc = Document(io.BytesIO(file_bytes))
        fresh_units = extract_units(doc)  # fresh copy bound to this doc
        count = apply_suggestions(fresh_units, [s for s in chosen if s["id"] in fresh_units])
        out = io.BytesIO()
        doc.save(out)
        st.session_state.output = out.getvalue()
        st.session_state.applied = count

    if st.session_state.get("output"):
        st.success(f"Applied {st.session_state.applied} edit(s).")
        st.download_button(
            "⬇️ Download edited Word file",
            data=st.session_state.output,
            file_name=uploaded.name.replace(".docx", "") + "_edited.docx",
            mime=DOCX_MIME,
        )

st.divider()
st.caption("Uploaded files are processed in memory and are not stored by this app.")
